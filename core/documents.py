import os
import re

import PyPDF2

from docx import Document


def clean_text(text):
    if not text:
        return ""

    text = str(text)

    text = text.replace("\x00", " ")

    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text
    )

    return text.strip()


def extract_pdf(path):

    text_parts = []

    try:

        with open(
            path,
            "rb"
        ) as f:

            reader = PyPDF2.PdfReader(f)

            for page in reader.pages:

                try:
                    text = page.extract_text() or ""

                except Exception:
                    text = ""

                if text.strip():
                    text_parts.append(text)

    except Exception as e:

        print(
            f"PDF extraction error "
            f"({os.path.basename(path)}): {e}"
        )

        return ""

    return clean_text(
        "\n\n".join(text_parts)
    )


def extract_docx(path):

    try:

        document = Document(path)

        parts = []

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                parts.append(text)

        # Also extract table content.
        for table in document.tables:

            for row in table.rows:

                cells = []

                for cell in row.cells:

                    value = cell.text.strip()

                    if value:
                        cells.append(value)

                if cells:
                    parts.append(
                        " | ".join(cells)
                    )

        return clean_text(
            "\n\n".join(parts)
        )

    except Exception as e:

        print(
            f"DOCX extraction error "
            f"({os.path.basename(path)}): {e}"
        )

        return ""


def extract_txt(path):

    encodings = [
        "utf-8",
        "utf-8-sig",
        "cp1252",
        "latin-1"
    ]

    for encoding in encodings:

        try:

            with open(
                path,
                "r",
                encoding=encoding
            ) as f:

                return clean_text(
                    f.read()
                )

        except UnicodeDecodeError:
            continue

        except Exception as e:

            print(
                f"TXT extraction error "
                f"({os.path.basename(path)}): {e}"
            )

            return ""

    return ""


def create_chunks(
    text,
    source_name="",
    chunk_size=900,
    overlap=150
):

    text = clean_text(text)

    if not text:
        return []

    if chunk_size <= 0:
        chunk_size = 900

    if overlap < 0:
        overlap = 0

    if overlap >= chunk_size:
        overlap = chunk_size // 5

    chunks = []

    start = 0
    text_length = len(text)

    while start < text_length:

        end = min(
            start + chunk_size,
            text_length
        )

        chunk_text = text[start:end].strip()

        if chunk_text:

            chunks.append(
                {
                    "text": chunk_text,
                    "source": source_name,
                }
            )

        if end >= text_length:
            break

        next_start = end - overlap

        if next_start <= start:
            next_start = end

        start = next_start

    return chunks


def extract_text_from_file(path):

    if not path:
        return ""

    extension = (
        os.path.splitext(path)[1]
        .lower()
    )

    if extension == ".pdf":
        return extract_pdf(path)

    if extension == ".docx":
        return extract_docx(path)

    if extension in {
        ".txt",
        ".md",
        ".csv",
        ".log"
    }:
        return extract_txt(path)

    return ""


def process_document(
    path,
    chunk_size=900,
    overlap=150
):

    text = extract_text_from_file(path)

    if not text:
        return []

    source_name = os.path.basename(path)

    return create_chunks(
        text=text,
        source_name=source_name,
        chunk_size=chunk_size,
        overlap=overlap
    )
