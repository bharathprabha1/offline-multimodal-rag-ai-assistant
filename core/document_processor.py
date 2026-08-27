import re
import os
import json
import csv
import io

try:
    import PyPDF2
    from docx import Document
except ImportError as e:
    raise ImportError(f"Document processing dependency missing: {e}")

try:
    from PIL import Image
    import pytesseract
except ImportError:
    Image = None
    pytesseract = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    from bs4 import BeautifulSoup
except ImportError:
    BeautifulSoup = None

try:
    import vosk
    import wave
except ImportError:
    vosk = None
    wave = None


# ============================================================
# TESSERACT
# ============================================================

TESSERACT_PATH = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

if pytesseract is not None:
    if os.path.exists(TESSERACT_PATH):
        pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH


# ============================================================
# TEXT CLEANING
# ============================================================

def clean_text(text):

    if not text:
        return ""

    text = str(text)

    text = text.replace("\x00", " ")

    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    text = re.sub(
        r"\n\s*\n+",
        "\n\n",
        text
    )

    return text.strip()


# ============================================================
# PDF
# ============================================================

def extract_pdf(path):

    text = ""

    try:

        with open(path, "rb") as f:

            reader = PyPDF2.PdfReader(f)

            pages = []

            for page in reader.pages:

                try:

                    page_text = page.extract_text()

                    if page_text:
                        pages.append(page_text)

                except Exception:
                    continue

            text = "\n".join(pages)

    except Exception as e:

        print(
            f"PDF error {path}:",
            e
        )

    return clean_text(text)


# ============================================================
# DOCX
# ============================================================

def extract_docx(path):

    text = ""

    try:

        document = Document(path)

        parts = []

        for paragraph in document.paragraphs:

            if paragraph.text.strip():
                parts.append(
                    paragraph.text
                )

        # Also extract table content
        for table in document.tables:

            for row in table.rows:

                values = []

                for cell in row.cells:
                    values.append(cell.text)

                parts.append(
                    " | ".join(values)
                )

        text = "\n".join(parts)

    except Exception as e:

        print(
            f"DOCX error {path}:",
            e
        )

    return clean_text(text)


# ============================================================
# TXT / MD
# ============================================================

def extract_txt(path):

    try:

        with open(
            path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as f:

            return clean_text(
                f.read()
            )

    except Exception as e:

        print(
            f"TXT error {path}:",
            e
        )

        return ""


# ============================================================
# CSV
# ============================================================

def extract_csv(path):

    try:

        rows = []

        with open(
            path,
            "r",
            encoding="utf-8",
            errors="ignore",
            newline=""
        ) as f:

            reader = csv.reader(f)

            for row in reader:

                values = [
                    str(value).strip()
                    for value in row
                    if str(value).strip()
                ]

                if values:
                    rows.append(
                        " | ".join(values)
                    )

        return clean_text(
            "\n".join(rows)
        )

    except Exception as e:

        print(
            f"CSV error {path}:",
            e
        )

        return ""


# ============================================================
# XLSX
# ============================================================

def extract_xlsx(path):

    if openpyxl is None:
        return ""

    try:

        workbook = openpyxl.load_workbook(
            path,
            read_only=True,
            data_only=True
        )

        lines = []

        for sheet in workbook.worksheets:

            lines.append(
                f"Sheet: {sheet.title}"
            )

            for row in sheet.iter_rows(
                values_only=True
            ):

                values = []

                for value in row:

                    if value is not None:
                        values.append(
                            str(value).strip()
                        )

                if values:
                    lines.append(
                        " | ".join(values)
                    )

        workbook.close()

        return clean_text(
            "\n".join(lines)
        )

    except Exception as e:

        print(
            f"XLSX error {path}:",
            e
        )

        return ""


# ============================================================
# PPTX
# ============================================================

def extract_pptx(path):

    if Presentation is None:
        return ""

    try:

        presentation = Presentation(path)

        slides = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1
        ):

            parts = [
                f"Slide {slide_number}:"
            ]

            for shape in slide.shapes:

                if hasattr(
                    shape,
                    "text"
                ):

                    value = shape.text.strip()

                    if value:
                        parts.append(value)

            if len(parts) > 1:
                slides.append(
                    "\n".join(parts)
                )

        return clean_text(
            "\n\n".join(slides)
        )

    except Exception as e:

        print(
            f"PPTX error {path}:",
            e
        )

        return ""


# ============================================================
# HTML
# ============================================================

def extract_html(path):

    if BeautifulSoup is None:
        return ""

    try:

        with open(
            path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as f:

            html = f.read()

        soup = BeautifulSoup(
            html,
            "html.parser"
        )

        for tag in soup(
            [
                "script",
                "style",
                "noscript"
            ]
        ):

            tag.decompose()

        text = soup.get_text(
            "\n"
        )

        return clean_text(text)

    except Exception as e:

        print(
            f"HTML error {path}:",
            e
        )

        return ""


# ============================================================
# JSON
# ============================================================

def extract_json(path):

    try:

        with open(
            path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as f:

            data = json.load(f)

        return clean_text(
            json.dumps(
                data,
                indent=2,
                ensure_ascii=False
            )
        )

    except Exception as e:

        print(
            f"JSON error {path}:",
            e
        )

        return ""


# ============================================================
# IMAGE OCR
# ============================================================

def extract_image(path):

    if Image is None or pytesseract is None:
        return ""

    try:

        image = Image.open(path)

        # Convert formats such as RGBA/P to RGB
        if image.mode not in (
            "RGB",
            "L"
        ):

            image = image.convert(
                "RGB"
            )

        text = pytesseract.image_to_string(
            image,
            config="--psm 6"
        )

        return clean_text(text)

    except Exception as e:

        print(
            f"OCR error {path}:",
            e
        )

        return ""


# ============================================================
# WAV / VOSK
# ============================================================

def process_audio(path):

    if vosk is None or wave is None:
        return ""

    try:

        wf = wave.open(
            path,
            "rb"
        )

        if wf.getnchannels() != 1:
            print(
                "Audio must be mono WAV:",
                path
            )
            wf.close()
            return ""

        model_path = os.environ.get(
            "VOSK_MODEL_PATH",
            ""
        )

        if not model_path or not os.path.isdir(
            model_path
        ):

            print(
                "VOSK_MODEL_PATH is not configured."
            )

            wf.close()
            return ""

        model = vosk.Model(
            model_path
        )

        recognizer = vosk.KaldiRecognizer(
            model,
            wf.getframerate()
        )

        parts = []

        while True:

            data = wf.readframes(
                4000
            )

            if not data:
                break

            if recognizer.AcceptWaveform(
                data
            ):

                result = json.loads(
                    recognizer.Result()
                )

                value = result.get(
                    "text",
                    ""
                )

                if value:
                    parts.append(value)

        final_result = json.loads(
            recognizer.FinalResult()
        )

        value = final_result.get(
            "text",
            ""
        )

        if value:
            parts.append(value)

        wf.close()

        return clean_text(
            " ".join(parts)
        )

    except Exception as e:

        print(
            f"Audio error {path}:",
            e
        )

        return ""


# ============================================================
# GENERIC DOCUMENT EXTRACTION
# ============================================================

def extract_document(path):

    if not path:
        return ""

    extension = os.path.splitext(
        path
    )[1].lower()

    if extension == ".pdf":
        return extract_pdf(path)

    if extension == ".docx":
        return extract_docx(path)

    if extension in (
        ".txt",
        ".md"
    ):
        return extract_txt(path)

    if extension == ".csv":
        return extract_csv(path)

    if extension == ".xlsx":
        return extract_xlsx(path)

    if extension == ".pptx":
        return extract_pptx(path)

    if extension in (
        ".html",
        ".htm"
    ):
        return extract_html(path)

    if extension == ".json":
        return extract_json(path)

    if extension in (
        ".jpg",
        ".jpeg",
        ".png",
        ".bmp",
        ".tif",
        ".tiff",
        ".webp"
    ):
        return extract_image(path)

    if extension == ".wav":
        return process_audio(path)

    return ""


# ============================================================
# CHUNKING
# ============================================================

def create_chunks(
    text,
    source,
    chunk_size=500,
    overlap=80
):

    words = text.split()

    if not words:
        return []

    result = []

    start = 0

    while start < len(words):

        end = min(
            start + chunk_size,
            len(words)
        )

        chunk_text = " ".join(
            words[start:end]
        )

        if len(chunk_text) > 30:

            result.append(
                {
                    "source": source,
                    "text": chunk_text
                }
            )

        if end >= len(words):
            break

        start = end - overlap

    return result


# ============================================================
# COMPATIBILITY
# ============================================================

def process_document(path):

    text = extract_document(path)

    return clean_text(text)
